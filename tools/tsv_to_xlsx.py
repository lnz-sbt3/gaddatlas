#!/usr/bin/env python3
"""
tsv_to_xlsx.py — Rigenera gli XLSX di data/source/xlsx/ dai TSV canonici (D-019).

I TSV in data/source/tables/ sono la sorgente: gli XLSX sono un derivato, una
superficie di editing per chi lavora in Excel. Chi modifica un XLSX deve
riportare la modifica nel TSV (python3 tools/xlsx_to_tsv.py, che sovrascrive
il TSV): altrimenti il build successivo la ignora.

Fedeltà del giro TSV -> XLSX -> TSV:
  - diventano numeri solo gli interi che clean() di xlsx_to_tsv.py riconverte
    nella stessa stringa ("12" sì, "012" no). I decimali restano testo:
    openpyxl li scrive con 16 cifre significative e le coordinate
    perderebbero l'ultima cifra al ritorno; "1.0" tornerebbe "1";
  - celle vuote -> celle vuote; testo multiriga conservato.

Uso:
    python3 tools/tsv_to_xlsx.py                 (tutti i TSV -> data/source/xlsx/)
    python3 tools/tsv_to_xlsx.py --tables DIR --out DIR
"""
import argparse, csv, re, sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xlsx_to_tsv import clean  # stessa regola di normalizzazione del giro inverso

INTEGER = re.compile(r"-?\d+")


def cell_value(text):
    if text == "":
        return None
    if INTEGER.fullmatch(text) and clean(int(text)) == text:
        return int(text)
    return text


def convert(tsv_path: Path, out_dir: Path) -> Path:
    with open(tsv_path, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f, delimiter="\t", quotechar='"'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tsv_path.stem
    for row in rows:
        ws.append([cell_value(c) for c in row])
    out = out_dir / (tsv_path.stem + ".xlsx")
    wb.save(out)
    print(f"OK  {tsv_path} -> {out}  ({len(rows) - 1} righe)")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--tables", default="data/source/tables")
    ap.add_argument("--out", default="data/source/xlsx")
    a = ap.parse_args()
    out_dir = Path(a.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    tsvs = sorted(Path(a.tables).glob("*.tsv"))
    if not tsvs:
        sys.exit(f"Nessun TSV in {a.tables}")
    for p in tsvs:
        convert(p, out_dir)


if __name__ == "__main__":
    main()
